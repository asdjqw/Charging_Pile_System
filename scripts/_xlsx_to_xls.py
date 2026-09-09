# -*- coding: utf-8 -*-
"""Copy the generated xlsx into a .xls that matches the user-client workbook look."""
from openpyxl import load_workbook
import xlwt

XLSX = r"E:\Learning Resourse\Small_s3\Charge_pile\管理员客户端功能测试用例.xlsx"
XLS = r"E:\Learning Resourse\Small_s3\Charge_pile\管理员客户端功能测试用例.xls"

FONT_NAME = "宋体"


def xf(bold=False, center=False, header=False):
    style = xlwt.XFStyle()
    font = xlwt.Font()
    font.name = FONT_NAME
    font.height = 20 * 10
    font.bold = bold
    style.font = font
    al = xlwt.Alignment()
    al.wrap = 1
    al.vert = xlwt.Alignment.VERT_CENTER
    al.horz = xlwt.Alignment.HORZ_CENTER if center else xlwt.Alignment.HORZ_LEFT
    style.alignment = al
    borders = xlwt.Borders()
    thin = xlwt.Borders.THIN
    borders.left = borders.right = borders.top = borders.bottom = thin
    style.borders = borders
    if header:
        pattern = xlwt.Pattern()
        pattern.pattern = xlwt.Pattern.SOLID_PATTERN
        pattern.pattern_fore_colour = 22  # silver/gray
        style.pattern = pattern
    style.num_format_str = "0"
    return style


def main():
    src = load_workbook(XLSX)
    wb = xlwt.Workbook(encoding="utf-8")
    label = xf(bold=True, center=True)
    value = xf(bold=False, center=False)
    header = xf(bold=True, center=True, header=True)
    id_style = xf(bold=False, center=True)
    for name in src.sheetnames:
        ws_src = src[name]
        ws = wb.add_sheet(name[:31])
        widths = [2560, 9216, 7168, 9216, 9216, 3072, 3584]
        for i, w in enumerate(widths):
            ws.col(i).width = w
        # 合并区只写左上角，避免 xlwt overwrite
        merges = [
            (0, 0, 1, 2), (0, 0, 4, 6),
            (1, 1, 1, 6),
            (2, 2, 1, 2), (2, 2, 4, 6),
            (3, 3, 1, 6), (4, 4, 1, 6), (5, 5, 1, 6),
        ]
        covered = set()
        for r1, r2, c1, c2 in merges:
            for rr in range(r1, r2 + 1):
                for cc in range(c1, c2 + 1):
                    if (rr, cc) != (r1, c1):
                        covered.add((rr, cc))

        heights = [400, 400, 400, 520, 520, 520, 400]
        for r in range(1, ws_src.max_row + 1):
            xl_r = r - 1
            ws.row(xl_r).height = heights[xl_r] if xl_r < 7 else 900
            for c in range(1, 8):
                if (xl_r, c - 1) in covered:
                    continue
                cell = ws_src.cell(r, c)
                val = cell.value
                if val is None:
                    val = ""
                if r <= 6:
                    is_label = bool(val) and ((c == 1) or (c == 4 and r in (1, 3)))
                    style = label if is_label else value
                elif r == 7:
                    style = header
                else:
                    style = id_style if c in (1, 6) else value
                    if c == 1 and val != "":
                        val = int(val)
                ws.write(xl_r, c - 1, val, style)
        for r1, r2, c1, c2 in merges:
            ws.merge(r1, r2, c1, c2)
    out = XLS
    try:
        wb.save(out)
    except PermissionError:
        out = r"E:\Learning Resourse\Small_s3\Charge_pile\管理员客户端功能测试用例_样式对齐.xls"
        wb.save(out)
    print("saved", out)


if __name__ == "__main__":
    main()
